# Story input, read only (HTTP GET); credentials from JIRA_URL / JIRA_EMAIL / JIRA_API_TOKEN in .env: Jira title / description / comments / remote links / attachments, and the Confluence pages
# the story links to. Attachments are downloaded into the run folder (deleted by the run) within the per-story caps in config.
import html, os, pathlib, re
import requests
from . import config

S = config.load('scope')


def _env():
    return os.environ['JIRA_URL'].rstrip('/'), (os.environ['JIRA_EMAIL'], os.environ['JIRA_API_TOKEN'])


class Client:
    def __init__(self):
        self.base, self.auth = _env()

    def get(self, path_or_url, stream=False, **params):
        url = path_or_url if path_or_url.startswith('http') else self.base + path_or_url
        r = requests.get(url, auth=self.auth, params=params or None, timeout=120, stream=stream); r.raise_for_status()
        return r if stream else r.json()


def adf(n):
    """Atlassian document format -> text (links kept as URLs), as the existing PI-12 fetch does."""
    if isinstance(n, dict):
        t = n.get('text', '') or ''
        if n.get('type') in ('inlineCard', 'blockCard'): t = ' ' + n.get('attrs', {}).get('url', '') + ' '
        if n.get('type') == 'mention': t = n.get('attrs', {}).get('text', '')
        if n.get('type') == 'text':
            for m in n.get('marks', []) or []:
                if m.get('type') == 'link' and m.get('attrs', {}).get('href'): t += f" ({m['attrs']['href']})"
        return t + ''.join(adf(c) for c in n.get('content', [])) + ('\n' if n.get('type') in ('paragraph', 'listItem', 'heading', 'tableRow') else ' ')
    return ''


def confluence_ids(text, base):
    """Confluence page ids linked in text on this Atlassian site (pages/<id> and pageId=<id> links)."""
    host = re.escape(re.sub(r'^https?://', '', base))
    ids = []
    for m in re.finditer(rf'https?://{host}/wiki/[^\s)\]>"\']+', text):
        u = m.group(0)
        i = re.search(r'/pages/(\d+)', u) or re.search(r'[?&]pageId=(\d+)', u)
        if i and i.group(1) not in ids: ids.append(i.group(1))
    return ids


def select_attachments(atts, caps):
    """Split attachment metadata into (read, not_read) by type and the per-story caps. Oldest first."""
    kinds = {e: 'image' for e in caps['images']} | {e: 'text' for e in caps['text']}
    read, skip, total = [], [], 0
    limit = caps['max_total_mb'] * 1024 * 1024
    for a in sorted(atts, key=lambda a: a.get('created', '')):
        ext = pathlib.Path(a['filename']).suffix.lower()
        if ext not in kinds: skip.append(dict(a, reason=f'type {ext or "unknown"} not read')); continue
        if len(read) >= caps['max_files']: skip.append(dict(a, reason=f"over the {caps['max_files']}-file cap")); continue
        if total + a.get('size', 0) > limit: skip.append(dict(a, reason=f"over the {caps['max_total_mb']} MB cap")); continue
        read.append(dict(a, kind=kinds[ext])); total += a.get('size', 0)
    return read, skip


def extract_text(p, max_pages):
    """Text of a document attachment, at most max_pages pages / slides. Returns (text, note)."""
    ext = p.suffix.lower()
    if ext in ('.md', '.txt'): return p.read_text(errors='ignore'), ''
    if ext == '.pdf':
        import pymupdf
        with pymupdf.open(p) as d:
            n = min(len(d), max_pages)
            return '\n'.join(d[i].get_text() for i in range(n)), (f'first {n} of {len(d)} pages' if len(d) > n else '')
    if ext == '.docx':
        import docx
        return '\n'.join(x.text for x in docx.Document(p).paragraphs), ''
    if ext == '.pptx':
        import pptx
        sl = list(pptx.Presentation(p).slides); n = min(len(sl), max_pages)
        out = []
        for i, s in enumerate(sl[:n], 1):
            out.append(f'[slide {i}] ' + ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
        return '\n'.join(out), (f'first {n} of {len(sl)} slides' if len(sl) > n else '')
    return '', 'not read'


def html_text(s):
    s = re.sub(r'<(br|/p|/li|/h\d|/tr)[^>]*>', '\n', s)
    return html.unescape(re.sub(r'<[^>]+>', ' ', s))


def story(key, inputs_dir):
    """Returns (title, text, images, inputs) for one story. text = description + COMMENTS + ATTACHMENTS + LINKED PAGES."""
    c = Client(); caps = S['inputs']['attachments']
    f = c.get('/rest/api/3/search/jql', jql=f'key = {key}', fields=','.join(S['jira']['fields']))['issues'][0]['fields']
    title, desc = f['summary'].strip(), adf(f.get('description'))
    comments, start = [], 0
    while True:
        d = c.get(f'/rest/api/3/issue/{key}/comment', startAt=start, maxResults=100, orderBy='created')
        comments += [dict(created=x['created'][:10], body=adf(x['body'])) for x in d['comments']]
        start += len(d['comments'])
        if start >= d['total'] or not d['comments']: break
    remote = [dict(title=x['object'].get('title'), url=x['object'].get('url')) for x in c.get(f'/rest/api/3/issue/{key}/remotelink')]
    text = desc + ('\n\nCOMMENTS\n' + '\n'.join(f"[comment {x['created']}] {x['body'].strip()}" for x in comments) if comments else '')
    # attachments
    atts = [dict(filename=a['filename'], size=a.get('size', 0), created=a.get('created', '')[:10], content=a['content'])
            for a in f.get('attachment') or []]
    read, not_read = select_attachments(atts, caps)
    inputs_dir.mkdir(parents=True, exist_ok=True)
    images, docs = [], []
    for i, a in enumerate(read):
        p = inputs_dir / f"{i:02d}-{re.sub(r'[^A-Za-z0-9._-]', '_', a['filename'])}"
        with c.get(a['content'], stream=True) as r, open(p, 'wb') as out:
            for chunk in r.iter_content(1 << 16): out.write(chunk)
        if a['kind'] == 'image': images.append(str(p)); continue
        try: t, note = extract_text(p, caps['max_pages'])
        except Exception as e: t, note = '', f'could not read: {type(e).__name__}'
        if not t.strip(): not_read.append(dict(a, reason=note or 'no text')); continue
        docs.append(dict(filename=a['filename'], created=a['created'], note=note, text=t.strip()))
    if docs:
        text += '\n\nATTACHMENTS\n' + '\n\n'.join(f"[attachment {d['filename']} {d['created']}{' – ' + d['note'] if d['note'] else ''}]\n{d['text']}" for d in docs)
    # linked Confluence pages (one hop)
    cp = S['inputs']['confluence']; wiki = c.base + S['jira']['confluence_path']
    ids = confluence_ids(text + '\n' + '\n'.join(x['url'] or '' for x in remote), c.base)
    pages, pages_not_read = [], []
    for pid in ids:
        if len(pages) >= cp['max_pages']: pages_not_read.append(dict(id=pid, reason=f"over the {cp['max_pages']}-page cap")); continue
        try:
            pg = c.get(f'{wiki}/rest/api/content/{pid}', expand='body.storage,version')
            pages.append(dict(id=pid, title=pg['title'], updated=pg['version']['when'][:10], text=html_text(pg['body']['storage']['value']).strip()))
        except Exception as e: pages_not_read.append(dict(id=pid, reason=f'could not read: {type(e).__name__}'))
    if pages:
        text += '\n\nLINKED CONFLUENCE PAGES\n' + '\n\n'.join(f"[page {p['title']} (id {p['id']}) updated {p['updated']}]\n{p['text']}" for p in pages)
    inputs = dict(comments=len(comments), remote_links=remote,
                  attachments_read=[dict(filename=d['filename'], note=d['note']) for d in docs] + [dict(filename=pathlib.Path(p).name, note='image') for p in images],
                  attachments_not_read=[dict(filename=a['filename'], reason=a['reason']) for a in not_read],
                  confluence_read=[dict(id=p['id'], title=p['title']) for p in pages], confluence_not_read=pages_not_read)
    return title, text, images, inputs
