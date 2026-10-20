# S3: story input – attachment types and caps, Confluence links (one hop), document text extraction.
from scope import config, jira

CAPS = config.load('scope')['inputs']['attachments']
MB = 1024 * 1024

def att(name, size=1000, created='2026-10-01'): return dict(filename=name, size=size, created=created, content='u')

def test_types_read_and_not_read():
    read, skip = jira.select_attachments([att('a.png'), att('b.pdf'), att('c.pptx'), att('d.xlsx'), att('e.ppt'), att('f.csv')], CAPS)
    assert [a['filename'] for a in read] == ['a.png', 'b.pdf', 'c.pptx']
    assert {a['filename'] for a in skip} == {'d.xlsx', 'e.ppt', 'f.csv'}

def test_file_and_size_caps():
    read, skip = jira.select_attachments([att(f'{i:02d}.txt', created=f'2026-10-{i + 1:02d}') for i in range(25)], CAPS)
    assert len(read) == 20 and len(skip) == 5 and all('20-file cap' in s['reason'] for s in skip)
    read, skip = jira.select_attachments([att('a.pdf', 20 * MB, '2026-01-01'), att('b.pdf', 6 * MB, '2026-01-02'), att('c.png', 4 * MB, '2026-01-03')], CAPS)
    assert [a['filename'] for a in read] == ['a.pdf', 'c.png'] and '25 MB' in skip[0]['reason']

def test_confluence_ids_same_site_only():
    base = 'https://wiki.example.test'
    t = ('see https://wiki.example.test/wiki/spaces/CE/pages/1111111111/PI-12+Preview and '
         'https://wiki.example.test/wiki/pages/viewpage.action?pageId=123 and https://other.example.test/wiki/spaces/X/pages/999 '
         'again https://wiki.example.test/wiki/spaces/CE/pages/1111111111')
    assert jira.confluence_ids(t, base) == ['1111111111', '123']

def test_adf_keeps_link_urls():
    n = {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'spec', 'marks': [{'type': 'link', 'attrs': {'href': 'https://h/wiki/x'}}]},
                                                                      {'type': 'inlineCard', 'attrs': {'url': 'https://h/wiki/y'}}]}]}
    t = jira.adf(n); assert 'https://h/wiki/x' in t and 'https://h/wiki/y' in t

def test_extract_text_docx_pptx_md_and_page_cap(tmp_path):
    import docx, pptx
    d = docx.Document(); d.add_paragraph('Fare rules for Payment'); d.save(tmp_path / 'a.docx')
    p = pptx.Presentation()
    for i in range(35):
        s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = f'slide text {i}'
    p.save(tmp_path / 'b.pptx'); (tmp_path / 'c.md').write_text('# md')
    assert 'Payment' in jira.extract_text(tmp_path / 'a.docx', 30)[0]
    t, note = jira.extract_text(tmp_path / 'b.pptx', 30)
    assert 'slide text 29' in t and 'slide text 30' not in t and note == 'first 30 of 35 slides'
    assert jira.extract_text(tmp_path / 'c.md', 30)[0] == '# md'
